from dataclasses import dataclass


@dataclass(slots=True)
class RawItem:
    row_number: int
    raw_text: str
    columns: dict[str, str]


NAME_KEYWORDS = (
    "наименование",
    "наимен-е",
    "наим.",
    "товар",
    "позиция",
    "название",
    "предмет",
    "name",
)
SKIP_PREFIXES = ("итого", "всего", "итог")


class ParserError(ValueError):
    pass


def _decode(data: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1251"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ParserError("Не удалось определить кодировку файла: ожидается UTF-8 или Windows-1251")


def _normalize(rows) -> list[list[str]]:
    return [[("" if value is None else str(value)).strip() for value in row] for row in rows]


def _read_csv_rows(data: bytes) -> list[list[str]]:
    import csv
    import io

    text = _decode(data)
    delimiter = ";"
    try:
        dialect = csv.Sniffer().sniff(text[:2048], delimiters=";,\t")
        delimiter = dialect.delimiter
    except csv.Error:
        pass
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    return _normalize(reader)


def _read_xlsx_rows(path) -> list[list[str]]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        return _normalize(sheet.iter_rows(values_only=True))
    finally:
        workbook.close()


def _header_index(rows: list[list[str]]) -> int:
    for index, row in enumerate(rows[:20]):
        non_empty = sum(1 for value in row if value)
        if non_empty >= 2:
            return index
    return 0


def _name_column(header: list[str], rows: list[list[str]]) -> int:
    for index, cell in enumerate(header):
        lowered = cell.lower()
        if lowered and any(keyword in lowered for keyword in NAME_KEYWORDS):
            return index
    width = max((len(row) for row in rows), default=0)
    best_column, best_total = 0, -1
    for column in range(width):
        total = sum(len(row[column]) for row in rows if len(row) > column)
        if total > best_total:
            best_column, best_total = column, total
    return best_column


def _columns(header: list[str], row: list[str]) -> dict[str, str]:
    result: dict[str, str] = {}
    for index, title in enumerate(header):
        key = title if title else f"колонка_{index + 1}"
        if key in result:
            key = f"{key}_{index + 1}"
        result[key] = row[index] if index < len(row) else ""
    return result


def parse_rows(rows: list[list[str]]) -> list[RawItem]:
    if not rows:
        raise ParserError("Файл пуст")
    header_index = _header_index(rows)
    header = rows[header_index]
    body = rows[header_index + 1 :]
    if not body:
        raise ParserError("В файле нет строк с позициями")
    name_column = _name_column(header, body)
    items: list[RawItem] = []
    for offset, row in enumerate(body):
        raw = row[name_column] if name_column < len(row) else ""
        if not raw or raw.lower().startswith(SKIP_PREFIXES):
            continue
        items.append(
            RawItem(
                row_number=header_index + offset + 2,
                raw_text=raw,
                columns=_columns(header, row),
            )
        )
    if not items:
        raise ParserError("Не найдено ни одной позиции: проверьте колонку с наименованием")
    return items


def parse_file(path) -> list[RawItem]:
    from pathlib import Path

    file_path = Path(path)
    extension = file_path.suffix.lower()
    if extension == ".csv":
        return parse_rows(_read_csv_rows(file_path.read_bytes()))
    if extension == ".xlsx":
        return parse_rows(_read_xlsx_rows(file_path))
    raise ParserError(f"Неподдерживаемый формат файла: {extension}")
