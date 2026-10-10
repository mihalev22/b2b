from openpyxl import load_workbook

from app.models.entities import Item
from app.services.exporter import items_to_csv, items_to_xlsx


def _item(raw_name: str) -> Item:
    return Item(
        row_number=2,
        raw_name=raw_name,
        raw_columns={},
        status="auto",
        confidence=90.0,
        ktru_code="26.20.11.110-00000001",
        ktru_name="Ноутбук",
    )


def test_xlsx_export_neutralizes_formula_injection(tmp_path):
    payload = "=HYPERLINK(\"http://evil\")"
    path = tmp_path / "out.xlsx"
    path.write_bytes(items_to_xlsx([_item(payload)]))
    sheet = load_workbook(path).active
    assert sheet.cell(row=2, column=2).value == f"'{payload}"


def test_csv_export_neutralizes_formula_injection():
    content = items_to_csv([_item("+1+1")]).decode("utf-8-sig")
    assert "'+1+1" in content


def test_catalog_search_escapes_wildcards():
    from app.api.catalog import like_pattern

    assert like_pattern("50%") == "%50\\%%"
    assert like_pattern("a_b") == "%a\\_b%"
    assert like_pattern("c:\\x") == "%c:\\\\x%"
