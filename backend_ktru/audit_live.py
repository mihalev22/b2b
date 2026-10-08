"""Воспроизводимая проверка жюри с НАСТОЯЩИМИ вызовами Ollama.
Создаёт синтетические файлы внутри audit_data; не утверждает подлинность тестовых кодов.
"""
import json
import math
import time
from pathlib import Path
import httpx
from PIL import Image, ImageDraw, ImageFont
from openpyxl import Workbook
from docx import Document
from pptx import Presentation
from pptx.util import Inches
import pymupdf
from step1_text_docs import parse_text_file
from step2_vision_images import parse_image_file
from step3_embeddings import to_vectors
from step6_confidence import judge

HERE = Path(__file__).parent
DATA = HERE / "audit_data"
DATA.mkdir(exist_ok=True)
SPEC = "Клавиотура Logitech K120; интерфейс USB; проводная; количество 5 шт."
report = {"started": time.strftime("%Y-%m-%dT%H:%M:%S"), "tests": []}


def save_report():
    (HERE / "audit_results.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")


def run(name, function, evaluate):
    start = time.perf_counter()
    try:
        result = function()
        entry = {"name": name, "seconds": round(time.perf_counter()-start, 2),
                 "checks": evaluate(result), "result": result}
    except Exception as error:
        entry = {"name": name, "seconds": round(time.perf_counter()-start, 2), "error": str(error)}
    report["tests"].append(entry)
    save_report()
    print(name, entry.get("checks", entry.get("error")), flush=True)


def check_product(goods):
    if not goods:
        return {"one_product": False}
    item = goods[0]
    attrs = json.dumps(item.get("attributes", []), ensure_ascii=False).lower()
    return {"one_product": len(goods)==1,
            "quantity_5": item.get("quantity")==5,
            "brand": item.get("brand", "").lower()=="logitech",
            "model": item.get("model", "").lower()=="k120",
            "usb_attribute": "usb" in attrs,
            "wired_attribute": "провод" in attrs or "wired" in attrs}


def make_fixtures():
    (DATA / "spec.txt").write_text(SPEC, encoding="utf-8")
    (DATA / "spec.csv").write_text("Наименование;Интерфейс;Подключение;Количество;Единица\nКлавиотура Logitech K120;USB;проводная;5;шт\n", encoding="utf-8")
    wb=Workbook(); wb.active.append(["Наименование", "Интерфейс", "Подключение", "Количество", "Единица"])
    wb.active.append(["Клавиотура Logitech K120", "USB", "проводная", 5, "шт"]); wb.save(DATA / "spec.xlsx")
    doc=Document(); doc.add_paragraph(SPEC); doc.save(DATA / "spec.docx")
    prs=Presentation(); slide=prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_textbox(Inches(0.5), Inches(1), Inches(9), Inches(3)).text=SPEC
    prs.save(DATA / "spec.pptx")
    font=ImageFont.truetype("C:/Windows/Fonts/arial.ttf", 32)
    image=Image.new("RGB", (1300, 300), "white")
    draw=ImageDraw.Draw(image)
    draw.text((30,40), "Клавиотура Logitech K120", font=font, fill="black")
    draw.text((30,100), "Интерфейс USB; проводная", font=font, fill="black")
    draw.text((30,160), "Количество 5 шт.", font=font, fill="black")
    for ext in ("png", "jpg", "jpeg"):
        image.save(DATA / ("spec." + ext))
    pdf=pymupdf.open(); page=pdf.new_page(width=900,height=300)
    page.insert_font(fontname="arial", fontfile="C:/Windows/Fonts/arial.ttf")
    page.insert_text((30,60), SPEC, fontname="arial", fontsize=14)
    pdf.save(DATA / "spec.pdf"); pdf.close()
    pdf=pymupdf.open(); page=pdf.new_page(width=900,height=300)
    page.insert_image(page.rect, filename=str(DATA / "spec.png")); pdf.save(DATA / "scan.pdf"); pdf.close()


def check_embeddings(vectors):
    sim=lambda a,b: sum(x*y for x,y in zip(a,b))
    return {"count": len(vectors), "dimension": len(vectors[0]),
            "finite": all(math.isfinite(n) for v in vectors for n in v),
            "unit_norm": all(abs(sum(n*n for n in v)-1)<1e-5 for v in vectors),
            "duplicate_similarity": sim(vectors[0],vectors[1]),
            "related_similarity": sim(vectors[0],vectors[2]),
            "unrelated_similarity": sim(vectors[0],vectors[3]),
            "semantic_order": sim(vectors[0],vectors[2])>sim(vectors[0],vectors[3])}


if __name__ == "__main__":
    make_fixtures()
    for ext in ("txt","csv","xlsx","docx","pptx","pdf"):
        run("document_"+ext, lambda ext=ext: parse_text_file(DATA / ("spec."+ext)), check_product)
    for ext in ("png","jpg","jpeg"):
        run("vision_"+ext, lambda ext=ext: parse_image_file(DATA / ("spec."+ext)), check_product)
    run("scan_pdf", lambda: parse_text_file(DATA / "scan.pdf"), check_product)
    run("embedding", lambda: to_vectors(["Клавиатура проводная USB", "Клавиатура проводная USB",
        "Клавиатура для компьютера с подключением USB", "Бумага офисная А4"]), check_embeddings)
    candidate={"code":"TEST-KEYBOARD", "name":"Клавиатура", "attributes":{"Интерфейс":"USB"}, "similarity":0.8}
    run("judge_match", lambda: judge({"name":"Клавиатура USB", "attributes":[{"name":"Интерфейс","value":"USB","unit":""}]},[candidate]),
        lambda r: {"selected_allowed":r["ktru_code"]=="TEST-KEYBOARD", "verdict":r["verdict"]})
    run("judge_wrong_type", lambda: judge({"name":"Бумага А4"},[candidate]),
        lambda r: {"abstained":r["ktru_code"] is None})
    run("judge_missing", lambda: judge({"name":"Клавиатура","attributes":[]},[candidate]),
        lambda r: {"not_auto_ok":r["verdict"]!="ok"})
    run("eis_http", lambda: {"status":httpx.get("https://zakupki.gov.ru/epz/ktru/search/results.html",timeout=20,follow_redirects=True).status_code},
        lambda r: {"http_200":r["status"]==200})
